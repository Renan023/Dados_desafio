from openpyxl.chart import BarChart, Reference
from openpyxl import load_workbook

def GraficoDashboard(aba_dash,Graficotitle, start_row=7):

    max_row = aba_dash.max_row
    
    categorias = Reference(aba_dash, min_col= 1, min_row= start_row + 1 , max_row= max_row)
    valores = Reference(aba_dash, min_col= 2 , min_row=start_row, max_row=max_row)
    
    chart = BarChart()
    chart.title = Graficotitle
    chart.add_data(valores,titles_from_data=True)
    chart.set_categories(categorias)
    aba_dash.add_chart(chart,"E5")